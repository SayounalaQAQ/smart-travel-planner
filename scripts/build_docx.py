#!/usr/bin/env python3
"""Build a user-facing trip handbook from a confirmed, structured travel plan."""
import argparse
import html
import json
import re
from pathlib import Path
from travel_utils import trip_readiness
from travel_readiness import apply_readiness


INTERNAL_TERMS=(r'\bMCP\b',r'tools/call',r'provider unavailable',r'health_check',r'acceptance session',
                r'\bretry\b',r'\bmock\b',r'\bfixture\b',r'\bdebug\b',r'\bsynthetic(?:\s+input)?\b',
                r'\binternal evidence\b',r'测试数据',r'测试环境',r'调试信息',r'内部证据',r'验收样例')
MAX_IMAGE_WIDTH=4.2
MAX_IMAGE_HEIGHT=2.65

def clarify_currency(data):
    """Resolve yen-shaped symbols only with explicit per-item currency, never geography."""
    def visit(value,context=None):
        if isinstance(value,str):
            return re.sub(r'[¥￥]\s*(?=\d)',(context+' ' if context in {'JPY','CNY'} else '币种未注明 '),value)
        if isinstance(value,list):return [visit(x,context) for x in value]
        if isinstance(value,dict):
            currency=value.get('currency') or value.get('price_currency') or context
            return {key:(item if key in {'source_url','page_url','image_url','local_path','currency','budget_currency','destination_currency'}
                         else visit(item,currency)) for key,item in value.items()}
        return value
    return visit(data)

def public_status(value):
    if ' + ' in str(value): return '、'.join(public_status(x) for x in str(value).split(' + '))
    return {'UNAVAILABLE':'暂未获取','UNVERIFIED':'待核验','PARTIAL':'部分完成','VERIFIED':'已核验',
        'CURRENT_REFERENCE':'当前参考，出发前复核','COMPOSITE_REFERENCE':'多段来源组合，完整时间待核',
        'ESTIMATED':'明确标注的估算','EXECUTION_PARTIAL':'执行资料部分完成','DATA_PARTIAL':'部分数据待补',
        'NOT_READY':'尚未就绪','READY':'已就绪','COMPLETE':'结构完整','PASS':'通过','FAIL':'未通过',
        'FORECAST_PENDING':'尚无目标日期可靠预报','HISTORICAL_REFERENCE':'历史气候参考',
        'BLOCKED':'存在待解决问题','NO_ROUTE_FOUND':'未找到可靠路线','DATA_PENDING':'待补资料',
        'UNKNOWN':'未知','NOT_APPLICABLE':'不适用','PER_PARTY':'全团','PER_PERSON':'每人','PER_ROOM':'每间房','PER_NIGHT':'每晚','PER_RIDE':'每次乘车',
        'VERIFIED_QUOTE':'有效目标日期报价','SLEEP_SCHEDULE_CONFLICT':'与晚起偏好冲突','WALKING_CONFLICT':'与步行上限冲突',
        'DUPLICATE_ACTIVITY':'重复活动','ROUTE_CLOSURE_RESEARCH':'进一步研究后补充路线',
        'PRICE_UNAVAILABLE':'目标价格暂未取得','PROPERTY_RESEARCH_COMPLETE':'酒店静态资料已取得'}.get(str(value),str(value))

def public_fact(value):
    return {'Varies by location':'各商户营业时间不同','Varies by store':'各店营业时间不同',
        'Open at all times':'全天开放','無料':'免费','たこ焼':'章鱼烧','ラーメン':'拉面'}.get(str(value),str(value))

def public_place(value):
    return {'AIRPORT_UNRESOLVED':'机场待确定','STATION_UNRESOLVED':'车站待确定',
            'HOTEL_UNRESOLVED':'住宿待确定','ACTIVITY_LOCATION_UNRESOLVED':'活动地点待确定'}.get(str(value),str(value))

def public_preference(value):
    return {'relaxed':'轻松','balanced':'平衡','intensive':'紧凑','late_start':'晚一些开始',
            'anime':'二次元','kimono':'和服','fireworks':'花火','traditional_culture':'传统文化',
            'figure_collecting':'手办收藏','retro_games':'复古游戏','pilgrimage':'作品巡礼'}.get(str(value),str(value))

def public_mode(value):
    return {'WALK':'步行','TRAIN':'铁路','METRO':'地铁','BUS':'巴士','TAXI':'出租车','FLIGHT':'飞行',
            'DRIVING':'驾车','TRANSIT':'公共交通','UNAVAILABLE':'方式待核验','UNKNOWN':'方式待核验'}.get(str(value),public_status(value))

def public_cost(value,currency,unknown=False):
    if value is None:return '未核实'
    if value==0 and unknown:return '尚无可计入的小计；其他费用未核实'
    return f'{currency} {value:,.2f}' if isinstance(value,(int,float)) else public_status(value)


def add_hyperlink(paragraph,label,url):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); props.append(color)
    under=OxmlElement('w:u'); under.set(qn('w:val'),'single'); props.append(under)
    run.append(props); node=OxmlElement('w:t'); node.text=label; run.append(node); link.append(run); paragraph._p.append(link)


def add_internal_link(paragraph,label,anchor):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),anchor)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); props.append(color)
    run.append(props); text=OxmlElement('w:t'); text.text=label; run.append(text); link.append(run); paragraph._p.append(link)


def add_bookmark(paragraph,name,bookmark_id):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(bookmark_id)); start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(bookmark_id))
    paragraph._p.insert(0,start); paragraph._p.append(end)


def add_section_heading(doc,text,level,anchor,bookmark_id):
    paragraph=doc.add_heading(text,level); add_bookmark(paragraph,anchor,bookmark_id); return paragraph


def set_cell_shading(cell,fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr=cell._tc.get_or_add_tcPr(); shd=tc_pr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tc_pr.append(shd)
    shd.set(qn('w:fill'),fill)


def format_table(table,header_rows=0):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table.autofit=True; borders=table._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is None: borders=OxmlElement('w:tblBorders'); table._tbl.tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node=borders.find(qn('w:'+edge))
        if node is None: node=OxmlElement('w:'+edge); borders.append(node)
        node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9')
    for row_index,row in enumerate(table.rows):
        tr_pr=row._tr.get_or_add_trPr()
        if row_index<header_rows:
            repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); tr_pr.append(repeat)
        for cell in row.cells:
            cell.vertical_alignment=1
            tc_pr=cell._tc.get_or_add_tcPr(); margins=tc_pr.find(qn('w:tcMar'))
            if margins is None: margins=OxmlElement('w:tcMar'); tc_pr.append(margins)
            for edge,amount in (('top','100'),('start','120'),('bottom','100'),('end','120')):
                node=margins.find(qn('w:'+edge))
                if node is None: node=OxmlElement('w:'+edge); margins.append(node)
                node.set(qn('w:w'),amount); node.set(qn('w:type'),'dxa')
        if row_index<header_rows:
            for cell in row.cells:
                set_cell_shading(cell,'F2F2F2')
                for p in cell.paragraphs:
                    for run in p.runs: run.bold=True


def visible_text(data):
    chunks=[]
    def visit(value):
        if isinstance(value,str): chunks.append(value)
        elif isinstance(value,dict):
            for key,item in value.items():
                if key == 'destination_choice_confirmed': continue
                if key not in {'evidence','field_provenance','raw_response','tool_response','route_provider','provider_debug',
                               'local_path','image_url','source_url','page_url','field_evidence'}: visit(item)
        elif isinstance(value,list):
            for item in value: visit(item)
    for key,value in data.items():
        if key == 'destination_choice_confirmed': continue
        if key not in {'evidence','field_provenance','raw_response','tool_response','provider_debug','mode','state_sync_ledger','execution_graph','cost_ledger','budget_reconciliation','route_closure','provider_probe_record','document_validation','hotel_search','flight_search','refresh_schedule','research_evidence','research_attempt_ledger','research_backfill_ledger','research_source_conflicts','cost_research_coverage','candidate_comparisons','research_decision_ledger','reference_budget','practical_plan_readiness','route_search_modes','restaurant_options','targeted_replanning_tasks','image_pipeline','experience_preferences'}: visit(value)
    return '\n'.join(chunks)


def sanitize_route_placeholders(value):
    """Remove legacy template wording without hiding real missing-data status."""
    if isinstance(value,str):
        value=re.sub(r'住宿(?:地址|地点).{0,28}(?:未定|尚未确定).{0,30}(?:选定|确定|确认)住宿后.{0,8}(?:补充|核验)', '住宿与门到门路线信息待核验', value)
        value=re.sub(r'(?:当前)?交通时间(?:为)?估算', '交通时间待核验', value)
        value=re.sub(r'门到门路线与分钟数(?:将)?待补', '门到门路线与分钟数待核验', value)
        return value
    if isinstance(value,list): return [sanitize_route_placeholders(item) for item in value]
    if isinstance(value,dict): return {key:sanitize_route_placeholders(item) for key,item in value.items()}
    return value


def validate_public_copy(data,technical_debug=False):
    if technical_debug: return []
    copy=visible_text(data); return [pattern for pattern in INTERNAL_TERMS if re.search(pattern,copy,re.I)]


def public_sources(data):
    entries=data.get('public_sources') if isinstance(data.get('public_sources'),list) else []
    seen=set(); clean=[]
    for item in entries:
        if not isinstance(item,dict): continue
        url=item.get('url') or item.get('source_url')
        if not isinstance(url,str) or not url.startswith('https://') or url in seen: continue
        title=str(item.get('title') or item.get('source_title') or '公开来源')
        if re.fullmatch(r'[a-z_]+',title):
            title={'overview':'日本旅行概览','fukuoka':'福冈目的地比较','tokyo':'东京目的地比较','areas':'大阪住宿区域',
                'hotel_index':'大阪酒店目录','hotel_guide':'住宿筛选参考','comfort':'Comfort 酒店官方资料','bridge':'The Bridge 酒店官方资料',
                'karaksa':'karaksa 酒店官方资料','flights':'上海至大阪航班参考','fuk_airport':'福冈机场接驳','kyoto_transport':'京都交通参考',
                'dotonbori':'道顿堀旅游资料','kuromon':'黑门市场旅游资料','park':'大阪城公园资料','castle':'大阪城天守阁资料',
                'tenshiba':'天王寺绿地资料','inari':'伏见稻荷资料','kiyomizu':'清水寺资料','gion':'祇园资料',
                'food_wanaka':'わなか餐饮资料','food_ajinoya':'味乃家餐饮资料','climate':'大阪历史气候','holidays':'日本节假日',
                'fx':'汇率参考','photo_dotonbori':'道顿堀图片许可','photo_castle':'大阪城图片许可','photo_inari':'伏见稻荷图片许可'}.get(title,'旅行公开参考')
        seen.add(url); clean.append({'title':title[:180],'url':url,'checked_at':str(item.get('checked_at') or '')[:40]})
    for attraction in data.get('attractions',[]):
        url=attraction.get('source_url')
        if isinstance(url,str) and url.startswith('https://') and url not in seen:
            seen.add(url); clean.append({'title':str(attraction.get('source_title') or attraction.get('name') or '景点官方来源')[:180],'url':url,'checked_at':str(attraction.get('source_checked_at') or '')[:40]})
    for image in data.get('images',[]):
        url=image.get('page_url')
        if isinstance(url,str) and url.startswith('https://') and url not in seen:
            seen.add(url); clean.append({'title':str(image.get('title') or image.get('source') or '图片来源')[:180],'url':url,'checked_at':str(image.get('retrieved_at') or '')[:40]})
    for item in data.get('research_evidence',[]):
        url=item.get('source_url')
        if url in seen:
            previous=next(c for c in clean if c['url']==url)
            if str(item.get('retrieved_at',''))>previous['checked_at']:
                previous.update(checked_at=str(item['retrieved_at'])[:40],title=str(item.get('source_title') or item.get('subject'))[:180])
        if isinstance(url,str) and url.startswith('https://') and url not in seen:
            seen.add(url); clean.append({'title':str(item.get('source_title') or item.get('subject'))[:180],
                'url':url,'checked_at':str(item.get('retrieved_at') or '')[:40]})
    return clean


def numeric_scenarios(data):
    budget=data.get('budget',{}); currency=budget.get('currency') or data.get('critical_info',{}).get('budget_currency') or ''
    scenarios=budget.get('scenarios') if isinstance(budget.get('scenarios'),dict) else {key:budget[key] for key in ('low','recommended','comfortable') if isinstance(budget.get(key),dict)}
    cards=[]
    for key,title in (('low','低预算'),('recommended','推荐'),('comfortable','舒适')):
        value=scenarios.get(key)
        if not isinstance(value,dict): continue
        total=value.get('total')
        if not isinstance(total,(int,float)):
            total=sum(float(v) for name,v in value.items() if name not in {'total','per_person'} and isinstance(v,(int,float)))
        count=budget.get('traveler_count') or data.get('trip_profile',{}).get('traveler_count') or data.get('critical_info',{}).get('traveler_count') or 1
        per_person=value.get('per_person')
        if not isinstance(per_person,(int,float)): per_person=total/max(1,int(count))
        cards.append({'title':title,'total':round(float(total),2),'per_person':round(float(per_person),2),'currency':currency})
    return cards,scenarios


def add_budget_cards(doc,data,technical=False):
    closure=data.get('budget_reconciliation')
    if isinstance(closure,dict):
        doc.add_paragraph('成本核算：已知小计不等于完整旅行总价；预算额度不计入成本。')
        table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('项目','全团','人均')): cell.text=label
        keys=(('user_budget_limit','用户预算上限'),('verified_cost','已核验成本小计'),('estimated_cost','参考/估算成本小计'))
        if technical:keys+=(('minimum_known_cost','已知成本下界'),('maximum_known_cost','已知成本上界'))
        for key,label in keys:
            if key!='user_budget_limit' and not technical and not closure.get(f'{key}_party'):continue
            row=table.add_row().cells; row[0].text=label
            for col,suffix in ((1,'party'),(2,'per_person')):
                value=closure.get(f'{key}_{suffix}'); row[col].text=public_cost(value,closure.get('currency',''),bool(closure.get('unknown_cost_categories')) and key!='user_budget_limit')
        row=table.add_row().cells; row[0].text='完整旅行总成本'
        for col,key in ((1,'party_total'),(2,'per_person_total')):
            value=closure.get(key); row[col].text='仍有未知成本，暂不能计算' if value is None else f"{closure.get('currency','')} {value:,.2f}"
        format_table(table,header_rows=1)
        categories={'flight':'机票','hotel':'住宿','transit':'交通','food':'餐饮','attractions':'景点','internet':'网络','insurance':'保险','visa_related':'签证相关','misc':'杂项','emergency':'备用金'}
        unknown='、'.join(categories.get(c,c) for c in closure.get('unknown_cost_categories',[]))
        if technical:doc.add_paragraph('预算可行性：'+public_status(closure.get('feasibility'))+'；未知项目：'+unknown)
        elif unknown:doc.add_paragraph('尚未核实完整价格的项目：'+unknown+'。取得这些报价后才能判断是否在预算内。')
        known_rows=[item for item in closure.get('rows',[]) if technical or item.get('amount') is not None]
        if not known_rows:return []
        table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('成本项目','单价/原币种','范围×数量','人均核算','证据状态')): cell.text=label
        for item in known_rows:
            row=table.add_row().cells; row[0].text=str(item.get('item',item.get('category')))
            if row[0].text==item.get('category'):row[0].text=categories.get(item['category'],row[0].text)
            row[1].text='未获得' if item.get('amount') is None else f"{item.get('currency','')} {item['amount']}"
            row[2].text=f"{public_status(item.get('scope',''))} × {item.get('quantity','?')}"
            row[3].text='未知' if item.get('per_person_total') is None else str(item['per_person_total'])
            row[4].text=public_status(item.get('cost_type'))
            source=item.get('source_url') or item.get('source')
            if source: add_hyperlink(row[4].paragraphs[0],'来源',source)
        format_table(table,header_rows=1)
        return []
    cards,scenarios=numeric_scenarios(data)
    if cards:
        table=doc.add_table(rows=2,cols=len(cards)); table.style='Light Shading Accent 1'
        for col,item in enumerate(cards):
            table.cell(0,col).text=item['title']; set_cell_shading(table.cell(0,col),'F2F2F2')
            table.cell(1,col).text=f"{item['currency']} {item['per_person']:,.0f} / 人\n全团 {item['currency']} {item['total']:,.0f}"
            for run in table.cell(0,col).paragraphs[0].runs: run.bold=True
        format_table(table,header_rows=1)
    currency=data.get('budget',{}).get('currency') or data.get('critical_info',{}).get('budget_currency') or ''
    cap_allocation=data.get('budget',{}).get('cap_allocation')
    if isinstance(cap_allocation,dict) and cap_allocation:
        doc.add_paragraph('规划上限分配（不是实时价格或成本可行性结论）')
        count=max(1,int(data.get('budget',{}).get('traveler_count') or 1))
        table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
        for cell,title in zip(table.rows[0].cells,('项目','全团上限分配','人均上限分配')): cell.text=title
        labels={'transport':'交通','accommodation':'住宿','attractions':'景点','food':'餐饮','shopping':'购物','internet':'网络','insurance':'保险','visa_related':'签证/入境','misc':'杂项','emergency':'应急金'}
        for key,value in cap_allocation.items():
            if not isinstance(value,(int,float)) or key in {'total','per_person'}: continue
            row=table.add_row().cells; row[0].text=labels.get(key,key); row[1].text=f'{currency} {value:,.0f}'; row[2].text=f'{currency} {value/count:,.0f}'
        format_table(table,header_rows=1)
    normalized=[]
    categories=('transport','accommodation','attractions','food','shopping','internet','insurance','visa_related','misc','emergency')
    labels={'transport':'交通','accommodation':'住宿','attractions':'景点','food':'餐饮','shopping':'购物','internet':'网络','insurance':'保险','visa_related':'签证/入境','misc':'杂项','emergency':'应急金'}
    for scenario,values in scenarios.items():
        if not isinstance(values,dict): continue
        for key in categories:
            if key in values: normalized.append((scenario,labels[key],values[key]))
    if normalized:
        table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
        for cell,title in zip(table.rows[0].cells,('方案','项目','金额')): cell.text=title
        for scenario,label,value in normalized:
            row=table.add_row().cells; row[0].text={'low':'低预算','recommended':'推荐','comfortable':'舒适'}.get(scenario,str(scenario)); row[1].text=label; row[2].text=f'{currency} {value}'
        format_table(table,header_rows=1)
    return cards


def picture_for_attraction(doc,cell,attraction,image,root,test_mode,image_root=None):
    if not image or image.get('license') in (None,'license_unknown'): return False
    local=Path(image.get('local_path','')).resolve(); allowed=(root/'outputs'/'images').resolve()
    if test_mode:
        test_assets=(root/'tests'/'fixtures').resolve(); accepted=(root/'acceptance').resolve()
        allowed_roots=(allowed,test_assets,accepted)
    else: allowed_roots=(allowed,Path(image_root).resolve()) if image_root else (allowed,)
    if not any(base==local or base in local.parents for base in allowed_roots) or not local.is_file(): return False
    from image_pipeline import attraction_id,LICENSE
    if not LICENSE.fullmatch(str(image.get('license',''))) or not image.get('page_url'):return False
    if image.get('attraction_id') and image['attraction_id']!=attraction_id(attraction):return False
    if not image.get('attraction_id') and image.get('attraction')!=attraction.get('name'):return False
    from PIL import Image
    from docx.shared import Inches
    Image.MAX_IMAGE_PIXELS=40_000_000
    with Image.open(local) as probe: probe.verify()
    with Image.open(local) as source:
        width,height=source.size; scale=min(MAX_IMAGE_WIDTH/width,MAX_IMAGE_HEIGHT/height)
        paragraph=cell.add_paragraph();paragraph.paragraph_format.keep_with_next=True
        run=paragraph.add_run();shape=run.add_picture(str(local),width=Inches(width*scale),height=Inches(height*scale))
        shape._inline.docPr.set('descr',attraction_id(attraction)+' | '+str(attraction.get('name')))
    paragraph.alignment=1
    caption=cell.add_paragraph(); caption.alignment=1; caption.paragraph_format.keep_together=True
    caption.add_run(f"{image.get('caption') or attraction.get('name')} · {html.unescape(str(image.get('author') or image.get('artist') or '作者未标明'))} · {image['license']} · ")
    add_hyperlink(caption,'图片来源',image['page_url'])
    license_url=image.get('license_url')
    if not license_url:
        match=re.fullmatch(r'CC (BY(?:-SA)?) (\d+(?:\.\d+)?)',str(image['license']),re.I)
        if match:license_url=f"https://creativecommons.org/licenses/{match[1].lower()}/{match[2]}/"
    if license_url and str(license_url).startswith('https://creativecommons.org/'):
        caption.add_run(' · ');add_hyperlink(caption,'许可条款',license_url)
    if image.get('transformations'):
        caption.add_run(' · 已校正方向并转换 JPEG，按比例缩放；图像沿用原许可。')
    return True


def add_attraction_card(doc,attraction,image,root,test_mode,bookmark_id,image_root=None):
    from docx.shared import Inches,Pt
    title=str(attraction.get('name') or '景点')
    heading=doc.add_heading(title,2); add_bookmark(heading,'attraction_'+str(bookmark_id),100+bookmark_id)
    if attraction.get('tagline'):
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; run=p.add_run(str(attraction['tagline'])); run.bold=True; run.font.size=Pt(10)
    if image and not picture_for_attraction(doc,doc,attraction,image,root,test_mode,image_root):
        doc.add_paragraph('配图未通过文件、许可或地点校验，已省略。')
    fields=(('为什么适合你','why_visit'),('最值得体验','main_highlights'),('与同类体验的区别','difference'),
            ('不适合情形','not_recommended_when'),('建议停留','estimated_visit_duration'),('价格参考','ticket_price'),
            ('开放时间','opening_hours'),('预约提醒','reservation_required'),('适合时段','best_time_of_day'),('交通','transport_info'),('不喜欢时的替代','alternative'))
    for label,key in fields:
        if attraction.get(key) is not None:
            p=doc.add_paragraph(); p.paragraph_format.keep_together=True; r=p.add_run(label+'：'); r.bold=True
            value=public_fact(attraction[key])
            if attraction.get('research_audited') and key in {'ticket_price','opening_hours','reservation_required','transport_info'}:
                e=attraction.get('field_evidence',{}).get(key)
                value=(public_fact(e['value'])+'（已读取来源中的当前信息；2027运营需复核）') if e else '尚未取得本次字段证据，出发前需确认'
            p.add_run(value)
    if attraction.get('source_url'):
        p=doc.add_paragraph(); p.paragraph_format.keep_together=True
        p.add_run(f"信息来源 · 查询于 {attraction.get('source_checked_at','未记录')} · ")
        add_hyperlink(p,'打开来源',attraction['source_url'])


def route_endpoint(value):
    """Presentation only: unknown anchors remain in the source execution graph."""
    text=str(value or '').strip()
    if not text or re.search(r'UNRESOLVED|未定|待定|待确定|待确认|尚未确定|^None$|^UNKNOWN$',text,re.I):
        return None
    return text


def day_heading(index,date):
    text=str(date or '').strip()
    if not text or re.match(r'^Day\s*\d+\b',text,re.I):
        return f'Day {index}（日期未定）'
    return f'Day {index} · {text}'


def route_summary(day,technical=False):
    chain=day.get('route_chain') or []
    if not chain:
        supplied=str(day.get('route_summary') or '')
        if re.search(r'住宿.{0,12}(?:确定|确认)后|(?:当前|交通).{0,16}(?:时间|分钟).{0,12}估算',supplied):
            return '路线信息待核验。'
        if not technical and (not supplied or 'UNRESOLVED' in supplied or '待确定 →' in supplied):
            return '接驳安排将在机场或具体住宿确定后细化；当天保留抵达、入住或返程缓冲。'
        return supplied or '路线信息待核验。'
    if not technical:
        places=[]
        for leg in chain:
            for key in ('from','to'):
                place=route_endpoint(leg.get(key))
                if place and (not places or places[-1]!=place):places.append(place)
        if len(places)<2:
            return '接驳安排将在机场或具体住宿确定后细化；当天保留抵达、入住或返程缓冲。'
        return ' → '.join(places)
    parts=[]
    for leg in chain:
        label=f"{public_place(leg.get('from'))} → {public_place(leg.get('to'))}"
        if isinstance(leg.get('duration_minutes'),(int,float)):
            label+=f"（约 {int(leg['duration_minutes'])} 分钟）"
        else: label+='（时间待核验）'
        if leg.get('mode'): label+=f" · {public_mode(leg['mode'])}"
        if isinstance(leg.get('transfers'),int): label+=f" · 换乘 {leg['transfers']} 次"
        if isinstance(leg.get('walking_minutes'),(int,float)): label+=f" · 步行 {int(leg['walking_minutes'])} 分钟"
        if leg.get('fare') is not None: label+=f" · 费用 {leg.get('currency','')} {leg['fare']}"
        parts.append(label)
    return '；'.join(parts)


def add_route_table(doc,day,technical=False):
    legs=day.get('route_chain') or []
    if not technical:
        legs=[leg for leg in legs if route_endpoint(leg.get('from')) and route_endpoint(leg.get('to'))
              and leg.get('from')!=leg.get('to')
              and (isinstance(leg.get('duration_minutes'),(int,float)) or leg.get('fare') is not None)]
    if not legs:
        return
    table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
    for cell,label in zip(table.rows[0].cells,('出发地 → 目的地','方式','预计时间/换乘','步行','费用')): cell.text=label
    for leg in legs:
        row=table.add_row().cells; row[0].text=f"{public_place(leg.get('from'))} → {public_place(leg.get('to'))}"; row[1].text=public_mode(leg.get('mode','待核验'))
        row[2].text=(f"约 {leg['duration_minutes']} 分钟" if isinstance(leg.get('duration_minutes'),(int,float)) else '完整时间待核验')+(f"；换乘 {leg['transfers']} 次" if isinstance(leg.get('transfers'),int) else '')
        if leg.get('status')=='COMPOSITE_REFERENCE':
            row[2].text+=f"；已知路段合计 {leg.get('known_component_minutes',0)} 分钟，不含候车、换乘、站内及未核步行"
        row[3].text=f"{leg.get('walking_minutes','待核验')} 分钟" if leg.get('walking_minutes') is not None else '待核验'
        row[4].text=f"{leg.get('currency','')} {leg.get('fare','待核验')}" if leg.get('fare') is not None else '待核验'
        row[4].text+='；'+public_status(leg.get('status','UNAVAILABLE'))
        source=leg.get('source_url') or leg.get('source')
        if source: add_hyperlink(row[4].paragraphs[0],'来源',source)
        for index,component in enumerate(leg.get('components',[]),1):
            add_hyperlink(row[4].paragraphs[0],f' 分段来源{index}',component['source_url'])
    format_table(table,header_rows=1)


def add_toc(doc,sections):
    for label,anchor in sections:
        p=doc.add_paragraph(style='List Number'); add_internal_link(p,label,'section_'+anchor)


def order_handbook_sections(doc,sections):
    """Move complete existing sections; keep cover, links and final section settings."""
    from docx.oxml.ns import qn
    body=doc._element.body;prefix=[];groups={};current=None
    for child in list(body):
        if child.tag==qn('w:sectPr'):continue
        names=[b.get(qn('w:name')) for b in child.iter(qn('w:bookmarkStart'))]
        anchor=next((n for n in names if n and n.startswith('section_')),None)
        if anchor:current=anchor;groups[current]=[]
        (groups[current] if current else prefix).append(child)
    for child in prefix+[c for _,anchor in sections for c in groups.get('section_'+anchor,[])]:
        body.remove(child)
        body.insert(len(body)-1 if body[-1].tag==qn('w:sectPr') else len(body),child)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--input',required=True); parser.add_argument('--output',default='outputs/Smart_Travel_Plan.docx'); parser.add_argument('--image-root'); parser.add_argument('--test-synthetic',action='store_true'); parser.add_argument('--technical-debug-report',action='store_true'); args=parser.parse_args()
    input_path=Path(args.input).resolve(); output=Path(args.output).resolve(); data=apply_readiness(sanitize_route_placeholders(json.loads(input_path.read_text(encoding='utf-8'))))
    root=Path(__file__).resolve().parents[1]
    test_mode=args.test_synthetic
    if test_mode:
        allowed_inputs=((root/'tests'/'fixtures').resolve(),(root/'acceptance'/'fixtures').resolve())
        if not any(base==input_path or base in input_path.parents for base in allowed_inputs): raise SystemExit('--test-synthetic input must live under tests/fixtures or acceptance/fixtures')
        if data.get('mode')!='TEST_SYNTHETIC': raise SystemExit('--test-synthetic requires mode=TEST_SYNTHETIC in the input plan')
    if data.get('mode')=='TEST_SYNTHETIC' and not test_mode: raise SystemExit('Test-only profile cannot enter the normal user document flow')
    state=trip_readiness(data)
    if not state['ready_for_formal_plan']:
        print(json.dumps({'status':'blocked','reason':'Incomplete trip profile: continue destination research and ask the user; do not create a formal itinerary or handbook.','trip_readiness':state},ensure_ascii=False,indent=2)); raise SystemExit(2)
    data=clarify_currency(data)
    if not args.technical_debug_report:
        banned=validate_public_copy(data)
        if banned: raise SystemExit('User-facing document contains internal development terms and was not generated: '+', '.join(banned))
    try:
        from docx import Document
        from docx.shared import Inches,Pt,RGBColor
    except ImportError as exc: raise SystemExit('DOCX support unavailable. Install the pinned optional dependency from requirements-docx.txt') from exc
    doc=Document(); section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.62); section.left_margin=section.right_margin=Inches(.72); section.different_first_page_header_footer=True
    styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
    from docx.oxml.ns import qn
    styles['Normal'].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    for name in ('Title','Heading 1','Heading 2','Heading 3'):
        styles[name].font.name='Aptos Display'; styles[name].font.color.rgb=RGBColor(0,0,0); styles[name].paragraph_format.keep_with_next=True; styles[name].paragraph_format.keep_together=True
        styles[name].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    h=doc.add_paragraph(style='Title'); h.alignment=1; h.paragraph_format.space_before=Pt(68); h.add_run(str(data.get('title') or '旅行计划'))
    p=doc.add_paragraph(); p.alignment=1; p.add_run(str(data.get('subtitle') or '旅行手册'))
    p=doc.add_paragraph(); p.alignment=1; p.add_run(str(data.get('summary') or '按已确认的旅行条件整理。'))
    doc.add_page_break()
    sections=[('旅行摘要','summary'),('旅行画像','profile'),('专属体验亮点','experiences'),('每日路线','daily'),('核心景点','attractions'),('特色活动专题','events'),('住宿区域','lodging'),('交通路线','transport'),('餐饮候选与排队备选','restaurants'),('预算方案','budget'),('省钱与天气备选','alternatives'),('本次主动舍弃','omitted'),('关键决策','decisions'),('计划修订记录','revisions'),('预约与行前清单','checklist'),('当前旅行准备度','readiness'),('来源索引','sources')]
    sections=[s for s in sections if s[1] not in {'attractions','omitted'} or data.get('attractions' if s[1]=='attractions' else 'omitted_attractions')]
    sections=[s for s in sections if s[1] not in {'experiences','events'} or data.get('experience_highlights' if s[1]=='experiences' else 'seasonal_events')]
    doc.add_heading('目录',1); add_toc(doc,sections); doc.add_page_break()
    bookmarks={name:index+1 for index,(_,name) in enumerate(sections)}
    add_section_heading(doc,'当前旅行准备度',1,'section_readiness',bookmarks['readiness'])
    readiness=data.get('readiness',{})
    if args.technical_debug_report:
        doc.add_paragraph('Travel Readiness Dashboard')
        status_labels=(('plan_status','计划状态'),('data_status','数据准备度'),('booking_readiness','预订准备度'),('execution_readiness','实际执行准备度'),('route_structure_status','路线结构'),('route_coverage','路线数据覆盖率'),('route_verification_status','路线核验'),('exact_date_flight_quote','目标日期往返机票'),('exact_date_hotel_quote','目标日期酒店'),('plan_qa','综合质量检查'))
        table=doc.add_table(rows=0,cols=2); table.style='Light Shading Accent 1'
        for key,label in status_labels:
            row=table.add_row().cells; row[0].text=label
            row[1].text=f"{readiness.get(key,0):.0%}" if key=='route_coverage' else public_status(readiness.get(key,'UNVERIFIED'))
        extra_statuses=(('预算可行性',(data.get('budget_reconciliation') or {}).get('feasibility','UNVERIFIED')),
                        ('实用行程准备度',(data.get('practical_plan_readiness') or {}).get('status','UNVERIFIED')),
                        ('用户约束检查',(data.get('execution_evaluation') or {}).get('constraint_qa','UNVERIFIED')),
                        ('天气预报状态',data.get('weather_data_status','FORECAST_PENDING')))
        for label,value in extra_statuses:
            row=table.add_row().cells; row[0].text=label; row[1].text=public_status(value)
        format_table(table)
        doc.add_heading('已确认 / 参考估算 / 未获得 / 需刷新',2)
        doc.add_paragraph('已确认：旅行日期、人数、用户预算上限和偏好。参考估算：当前交通与景点资料，按标注时间适用。未获得：以下数据矩阵标为不可用的项目。需刷新：机票报价6小时、酒店报价12小时、汇率24小时、路线7天；临近出发重新查天气。')
        document=data.get('document_validation',{})
        doc.add_paragraph('文档结构检查：'+public_status(document.get('structural_status','UNVERIFIED'))+'；逐页视觉与分页复核：'+public_status(document.get('visual_status','UNVERIFIED'))+'。')
        research=data.get('research_coverage',{})
        if research:
            doc.add_paragraph(f"本次已补齐 {research.get('resolved',0)} 项当前参考，另有 {research.get('partial',0)} 项只查到部分资料。未查证项目不视为不存在；查询达到上限后保留待核清单。")
        doc.add_page_break()
        doc.add_heading('待核数据与质量详情',1)
        if readiness.get('blocking_reasons'):
            doc.add_heading('还缺哪些条件',2)
            blocker_labels={'EXACT_FLIGHT_QUOTE_MISSING':'目标日期往返航班报价未取得','EXACT_HOTEL_QUOTE_MISSING':'目标日期酒店库存和住宿总价未取得','AIRPORT_UNRESOLVED':'出入境机场尚未确定','HOTEL_PROPERTY_UNRESOLVED':'具体酒店及地址尚未确定','BUDGET_FEASIBILITY_UNVERIFIED':'总成本未闭环，尚不能判断预算是否足够','CRITICAL_ROUTE_EDGES_UNRESOLVED':'机场/跨城等关键交通段未完成核验','ROUTE_VERIFICATION_INCOMPLETE':'部分每日移动路线仍待核验','ATTRACTION_OPERATIONS_UNVERIFIED':'部分景点运营/预约条件仍需核验','WALKING_LIMIT_CHECK_INCOMPLETE':'步数上限尚未完成逐日核验'}
            for code in readiness['blocking_reasons']:
                fallback='目标日期酒店报价渠道目前不可用' if code.startswith('PROVIDER_HOTEL_QUOTE_') else '目标日期航班报价渠道目前不可用' if code.startswith('PROVIDER_FLIGHT_QUOTE_') else '仍有资料需要确认'
                doc.add_paragraph(blocker_labels.get(code,fallback),style='List Bullet')
        refresh=data.get('refresh_schedule',{})
        for key,label in (('next_quote_retry_after','下一次机酒核价检查点'),('route_review_by','出发前路线复核'),('weather_final_review_by','出发前天气复核')):
            if refresh.get(key): doc.add_paragraph(label+'：'+str(refresh[key]))
        matrix=data.get('data_completeness_matrix',[])
        if matrix:
            table=doc.add_table(rows=1,cols=2); table.style='Light Shading Accent 1'
            table.rows[0].cells[0].text='数据项'; table.rows[0].cells[1].text='状态'
            for item in matrix:
                row=table.add_row().cells; row[0].text=str(item.get('data','')); row[1].text=public_status(item.get('status','UNVERIFIED'))
            format_table(table,header_rows=1)
        checks=(data.get('plan_qa') or {}).get('checks',[])
        if checks:
            labels={'structural_qa':'行程结构','decision_qa':'决策一致性','constraint_qa':'用户约束','route_qa':'路线核验','budget_qa':'预算闭环','data_qa':'数据准备度','execution_qa':'可执行性','document_qa':'文档检查','user_preferences':'用户偏好','hard_constraints_and_walking':'硬约束与步行','budget_evidence':'预算证据','fatigue_and_step_estimate':'体力/步数估算','opening_hours':'开放时间','reservation':'预约','transport':'交通','geographic_clustering':'地理聚类','weather':'天气','attraction_diversity':'景点多样性','restaurant_diversity':'餐饮多样性','date_logic':'日期逻辑','arrival_departure':'到达/离开','lodging_location':'住宿位置','flight_impact':'航班影响','route_structure':'路线结构'}
            table=doc.add_table(rows=1,cols=2); table.style='Light Shading Accent 1'
            table.rows[0].cells[0].text='质量检查'; table.rows[0].cells[1].text='状态'
            for item in checks:
                row=table.add_row().cells; row[0].text=labels.get(item.get('check'),str(item.get('check'))); row[1].text=public_status(item.get('status','UNVERIFIED'))
            format_table(table,header_rows=1)
        findings=(data.get('plan_qa') or {}).get('findings',[])
        if findings:
            doc.add_paragraph('质量检查发现以下待处理项：')
            for item in findings:
                doc.add_paragraph(f"{item.get('code','PLAN_QA')}：{item.get('reason') or item.get('activity') or item.get('day') or '需要核验/重新规划'}",style='List Bullet')
        if readiness.get('booking_readiness') != 'READY':
            doc.add_paragraph('本手册提供的是规划结果；酒店库存、目标日期航班或交通等关键动态信息尚未全部核验，不能视为已经可预订。',style='Intense Quote')
    else:
        doc.add_paragraph('这份手册提供旅行建议，没有代为订票或订房。请先确认往返交通、具体酒店及住宿总价，再安排需要预约的体验。')
        doc.add_paragraph('未取得可靠资料的交通时间和费用不填数值；交通段的完整时间、站口与换乘请在住宿和机场确定后复核。临近出发再确认营业时间、天气和步行负荷。')
    add_section_heading(doc,'旅行摘要',1,'section_summary',bookmarks['summary']); doc.add_paragraph(str(data.get('summary') or '按已确认的旅行条件整理。'))
    add_section_heading(doc,'旅行画像',1,'section_profile',bookmarks['profile'])
    profile=data.get('trip_profile') or data.get('profile',{})
    if profile:
        table=doc.add_table(rows=0,cols=2); table.style='Light Shading Accent 1'
        profile_labels={'destination':'目的地','travel_window':'旅行日期','duration_days':'天数','traveler_count':'人数',
            'origin':'出发地','budget_amount':'用户预算上限','budget_currency':'预算币种','budget_scope':'预算口径',
            'airfare_included':'预算是否含机票','lodging_area':'住宿区域','hotel':'酒店定位候选（需核验库存）',
            'hotel_address':'候选酒店地址','origin_airport':'出发机场候选','arrival_airport':'到达机场候选',
            'intercity_mode':'往返方式','origin_station':'出发车站候选','arrival_station':'到达车站候选',
            'departure_station':'返程出发车站候选','return_station':'返程到达车站候选',
            'pace':'节奏','sleep_schedule':'作息偏好','walking_tolerance':'步行偏好','walking_tolerance_steps':'每日步数上限',
            'interests':'兴趣','food_preferences':'餐饮偏好','must_visit':'必去','avoid_list':'避开','visa_constraints':'入境条件'}
        for key,label in profile_labels.items():
            if key not in profile: continue
            if key=='airfare_included' and profile.get('intercity_mode')=='TRAIN':continue
            value=profile[key]
            if key=='budget_scope': value={'PER_PERSON':'每人','PER_PARTY':'全团'}.get(value,value)
            elif key=='airfare_included': value='包含' if value else '不包含'
            elif key=='intercity_mode': value=public_mode(value)
            elif key=='budget_currency':value={'CNY':'CNY 人民币','JPY':'JPY 日元'}.get(value,value)
            elif key in {'pace','sleep_schedule'}:value=public_preference(value)
            elif isinstance(value,list): value='、'.join(public_preference(x) for x in value)
            row=table.add_row().cells; row[0].text=label; row[1].text=str(value)
        format_table(table)
    add_section_heading(doc,'本次行程的关键决策',1,'section_decisions',bookmarks['decisions'])
    for item in data.get('key_decisions',[]): doc.add_paragraph(str(item),style='List Number')
    for item in data.get('decision_ledger',[]):
        if not isinstance(item,dict): continue
        doc.add_heading(str(item.get('category','Agent 决策')),2)
        doc.add_paragraph('选择：'+str(item.get('selected_option','未选择')))
        doc.add_paragraph('原因：'+str(item.get('why_selected','基于已核验候选和旅行约束推荐。')))
        if item.get('tradeoff'): doc.add_paragraph('取舍：'+str(item['tradeoff']))
        if item.get('confidence') and args.technical_debug_report: doc.add_paragraph('置信度：'+str(item['confidence']))
        if item.get('evidence'): doc.add_paragraph('依据：目的地、交通和商户公开资料，见来源索引；目标日期再次确认。')
    for item in data.get('user_decision_ledger',[]):
        if not isinstance(item,dict): continue
        doc.add_heading('需要你决定的取舍',2)
        doc.add_paragraph('问题：'+str(item.get('question','')))
        if item.get('reason'): doc.add_paragraph('原因：'+str(item['reason']))
        if item.get('impact'): doc.add_paragraph('对行程的影响：'+str(item['impact']))
        if item.get('default_recommendation'): doc.add_paragraph('默认建议：'+str(item['default_recommendation']))
    if data.get('price_priority_alternative'): doc.add_paragraph('如果你更看重价格，可以改成：'+str(data['price_priority_alternative']))
    if data.get('convenience_priority_alternative'): doc.add_paragraph('如果你更看重便利，可以改成：'+str(data['convenience_priority_alternative']))
    add_section_heading(doc,'为什么计划会这样变化',1,'section_revisions',bookmarks['revisions'])
    revisions=list(data.get('revision_ledger',[]))
    mode_labels={'WALK':'步行','TRAIN':'铁路组合','METRO':'地铁'}
    legacy_labels={'kyoto-dawn-higashiyama-plus-full-fushimi-hike':'清晨东山串游与伏见稻荷全山徒步',
                   'fushimi-late-morning-short-visit':'晚出发的伏见稻荷短访'}
    revisions=[{**rev,'old_decision':legacy_labels.get(rev.get('old_decision'),rev.get('old_decision')),
                'new_decision':legacy_labels.get(rev.get('new_decision'),rev.get('new_decision'))} for rev in revisions]
    for rev in data.get('decision_revision_ledger',[]):
        if rev.get('trigger')!='RESEARCH_DRIVEN_DECISION_CHANGE':continue
        old=rev['old_decision'];new=rev['new_decision'];route=new.get('route',new)
        revisions.append({'old_decision':mode_labels.get(old.get('mode'),old.get('place','原候选')),
            'new_decision':mode_labels.get(route.get('mode'),new.get('place','新候选')),
            'trigger':'新网页证据改变路线比较','reason':'比较已知步行负荷、换乘与未知环节后暂选；全天步行与连接时间仍待核验。',
            'evidence':['官方访问指引与旅游局路线参考，见来源索引']})
    if revisions:
        table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('原决定','新决定','触发冲突','改变原因','证据/影响')): cell.text=label
        for item in revisions:
            row=table.add_row().cells
            row[0].text=str(item.get('old_decision') or '未记录')
            row[1].text=str(item.get('new_decision') or '未记录')
            triggers=[item['trigger']] if item.get('trigger') else item.get('triggers',[])
            row[2].text='、'.join(public_status(x) for x in triggers)
            row[3].text=str(item.get('reason') or '')
            row[4].text='；'.join(map(str,item.get('evidence',[])))
        format_table(table,header_rows=1)
    else:
        doc.add_paragraph('本方案没有记录到需要推翻的重大决定。')
    if data.get('experience_highlights'):
        add_section_heading(doc,'用户专属体验亮点',1,'section_experiences',bookmarks['experiences'])
        for item in data['experience_highlights']:
            doc.add_paragraph(str(item),style='List Bullet')
    if data.get('seasonal_events'):
        add_section_heading(doc,'特色活动专题',1,'section_events',bookmarks['events'])
        for event in data['seasonal_events']:
            doc.add_heading(str(event.get('name','季节活动')),2)
            announced=event.get('status')=='OFFICIALLY_ANNOUNCED' and event.get('official_source_url') and event.get('date')
            doc.add_paragraph('日期：'+(str(event['date'])+'（官方公告参考，出发前复核）' if announced else '目标年份日期未核实；仅作季节候选，不保证参加，不作为出票依据。'))
            doc.add_paragraph(str(event.get('experience') or ''))
            doc.add_paragraph('替代：'+str(event.get('alternative') or '保留同区域自由活动，待官宣后再移动日程。'))
            if event.get('source_url'):add_hyperlink(doc.add_paragraph(),'活动官方来源',event['source_url'])
    add_section_heading(doc,'预算方案',1,'section_budget',bookmarks['budget'])
    doc.add_paragraph('币种：CNY 人民币；JPY 日元。用户预算是上限，不是必须花完的金额；报价、当前参考、规划估算与未知费用分别列示。')
    add_budget_cards(doc,data,technical=args.technical_debug_report)
    references=data.get('reference_budget',{}).get('reference_ledger',[])
    if references:
        doc.add_heading('已查到的价格参考',2)
        doc.add_paragraph('以下为公开菜单、单程票价或单个套餐参考，保留原币种。不是全程报价，也不表示已订到。目标日期往返交通和酒店总价仍未取得；用户预算仅为上限。')
        table=doc.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('项目','价格与口径','使用限制','来源')):cell.text=label
        scopes={'PER_ITEM':'每份/单品','PER_RIDE':'成人单程','PER_TICKET':'入场','PER_PACKAGE':'单个套餐'}
        for r in references:
            row=table.add_row().cells;row[0].text=str(r.get('subject',''))+' · '+r['item']
            row[1].text=f"{r['amount']:g} {r['currency']} / {scopes.get(r['scope'],'按页面口径')}"
            row[2].text='；'.join(r.get('limitations',[])) or '当前单项参考；目标日期复核'
            add_hyperlink(row[3].paragraphs[0],'价格来源',r['source']);row[3].add_paragraph(r['checked_at'][:10])
        format_table(table,header_rows=1)
        labels={'LOW_COST_SCENARIO':'低费用情景','BALANCED_SCENARIO':'均衡情景','COMFORT_SCENARIO':'舒适情景'}
        for key,scenario in data['reference_budget'].get('scenarios',{}).items():
            if scenario.get('total') is None:doc.add_paragraph(labels.get(key,'预算情景')+'：资料不足，暂不显示总额。')
            else:doc.add_paragraph(f"{labels.get(key,'预算情景')}：{scenario['currency']} {scenario['total']:,.2f} / 全团；人均 {scenario['per_person']:,.2f}。按明确数量假设与所列价格参考计算，非目标日期可订报价；未列项目与备用金另计。")
    budget=data.get('budget',{})
    budget_status=data.get('budget_status')
    if isinstance(budget_status,dict):
        table=doc.add_table(rows=0,cols=2); table.style='Light Shading Accent 1'
        labels=(('user_budget_limit','用户预算上限'),('current_verified_cost','已核验成本小计'),('current_estimated_cost','参考/估算成本小计'),('current_plan_cost','当前方案总成本'),('lowest_found_cost','已查到的最低成本'),('realistic_minimum_cost','最低合理成本'),('budget_gap','预算差额'),('feasibility','预算可行性'),('evidence_status','费用资料完成度'))
        for key,label in labels:
            row=table.add_row().cells; row[0].text=label
            value=budget_status.get(key)
            row[1].text=public_cost(value,data.get('trip_profile',{}).get('budget_currency',''),bool(data.get('budget_reconciliation',{}).get('unknown_cost_categories')) and key!='user_budget_limit')
        format_table(table)
        if budget_status.get('feasibility')=='UNVERIFIED':
            doc.add_paragraph('注意：预算仅为上限/规划估算，当前没有足够的目标日期价格证据证明可行。',style='Intense Quote')
    if budget.get('note'): doc.add_paragraph(str(budget['note']))
    if budget.get('optimization_summary'): doc.add_paragraph(str(budget['optimization_summary']))
    add_section_heading(doc,'住宿区域',1,'section_lodging',bookmarks['lodging'])
    areas=data.get('lodging_areas',[])
    if areas:
        table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('区域','交通','价格级别','覆盖','优缺点')): cell.text=label
        for area in areas:
            row=table.add_row().cells; row[0].text=str(area.get('name','')); row[1].text=str(area.get('transport','')); row[2].text=str(area.get('price_level',''))
            row[3].text=str(area.get('coverage','')); row[4].text=f"{area.get('pros','')}；{area.get('cons','')}"
        format_table(table,header_rows=1)
    if data.get('lodging_recommendation'): doc.add_paragraph(str(data['lodging_recommendation']))
    hotels=[h for h in data.get('hotel_availability',[]) if args.technical_debug_report or h.get('hotel_name')]
    if not args.technical_debug_report:
        for hotel in hotels:
            if hotel.get('total_price') is not None:continue
            doc.add_paragraph(str(hotel['hotel_name'])+'：住宿资料候选；'+str(hotel.get('area') or '')+'。目标日期房型、库存和住宿总价尚未核实，不视为已选定或已订到。')
            if hotel.get('source_url'):add_hyperlink(doc.add_paragraph(),'酒店官方资料',hotel['source_url'])
        hotels=[h for h in hotels if h.get('total_price') is not None]
    if hotels:
        table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('酒店候选','区域/房型/人数','目标日期','价格/价格类型','库存/来源')): cell.text=label
        for hotel in hotels:
            row=table.add_row().cells; row[0].text=str(hotel.get('hotel_name',''))
            row[1].text=f"{hotel.get('area','')}；{hotel.get('room_type') or '房型待核'}；{hotel.get('guest_count','?')}人"
            row[2].text=f"{hotel.get('check_in','')} 至 {hotel.get('check_out','')}"
            row[3].text=(f"{hotel.get('currency','')} {hotel.get('total_price')}" if hotel.get('total_price') is not None else '未获得')+'；'+public_status(hotel.get('price_type','UNAVAILABLE'))
            row[4].text=public_status(hotel.get('availability','UNKNOWN'))
            if hotel.get('source_url'): add_hyperlink(row[4].paragraphs[0],'来源',hotel['source_url'])
        format_table(table,header_rows=1)
    if data.get('hotel_ranking') and args.technical_debug_report:
        doc.add_paragraph('酒店比较状态：'+public_status(data['hotel_ranking'].get('status','DATA_PENDING')))
        if data['hotel_ranking'].get('reason'): doc.add_paragraph(str(data['hotel_ranking']['reason']))
    add_section_heading(doc,'交通路线',1,'section_transport',bookmarks['transport'])
    if data.get('transport_recommendation'): doc.add_paragraph(str(data['transport_recommendation']))
    transit_routes=[] if any(d.get('execution_graph') for d in data.get('days',[])) else data.get('transit_routes',[])
    if not args.technical_debug_report:
        transit_routes=[r for r in transit_routes if route_endpoint(r.get('from')) and route_endpoint(r.get('to'))
                        and r.get('from')!=r.get('to') and (r.get('duration_minutes') is not None or r.get('fare') is not None)]
    if transit_routes:
        table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('出发站 → 到达站','运营方/线路/换乘','耗时','票价/状态','来源')): cell.text=label
        for route in transit_routes:
            row=table.add_row().cells; row[0].text=str(route.get('from','')); row[1].text=str(route.get('to',''))
            row[1].text=f"{route.get('operator') or '待核验'} / {route.get('line') or '待核验'}；换乘：{'、'.join(route.get('transfer') or []) or '无/待确认'}"
            row[2].text=f"{route.get('duration_minutes')} 分钟" if route.get('duration_minutes') is not None else '待核验'
            row[3].text=(f"{route.get('currency','')} {route.get('fare')}" if route.get('fare') is not None else '待核验')+f"；{route.get('status','UNAVAILABLE')}"
            if route.get('source'): add_hyperlink(row[4].paragraphs[0],'打开来源',route['source'])
            else: row[4].text='暂无来源'
        format_table(table,header_rows=1)
    flights=[f for f in data.get('flight_availability',[]) if args.technical_debug_report or
             (route_endpoint(f.get('origin_airport')) and route_endpoint(f.get('destination_airport')))]
    if flights:
        table=doc.add_table(rows=1,cols=5); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('航线/日期','航司/航班/时间','行李','报价','状态/来源')): cell.text=label
        for flight in flights:
            row=table.add_row().cells
            row[0].text=f"{flight.get('origin_airport') or '机场待定'} → {flight.get('destination_airport') or '机场待定'}；{flight.get('flight_date','')}"
            row[1].text=f"{flight.get('airline') or '待核验'} {flight.get('flight_number') or ''}；{flight.get('departure_time') or '时刻待核'}–{flight.get('arrival_time') or '待核'}"
            row[2].text=str(flight.get('baggage') or '待核验')
            row[3].text=f"{flight.get('currency','')} {flight.get('price')}" if flight.get('price') is not None else '未获得'
            row[4].text=public_status(flight.get('status','UNAVAILABLE'))
            if flight.get('source'): add_hyperlink(row[4].paragraphs[0],'来源',flight['source'])
        format_table(table,header_rows=1)
    if not flights and not args.technical_debug_report:
        doc.add_paragraph('往返接驳：实际机场和航班确定后，再选择进城与返程交通；到达和离开当天留出手续、行李与换乘时间。当前不提供未经核实的时刻或票价。')
    if data.get('flight_ranking') and args.technical_debug_report:
        doc.add_paragraph('航班比较状态：'+public_status(data['flight_ranking'].get('status','DATA_PENDING'))+'；默认比较整体性价比，但目标日期报价不足时不选具体航班。')
    from image_pipeline import attraction_id
    images={item.get('attraction_id') or item.get('attraction'):item for item in data.get('images',[]) if isinstance(item,dict)}
    attractions=data.get('attractions',[])
    if attractions:
        add_section_heading(doc,'核心景点',1,'section_attractions',bookmarks['attractions'])
        for index,item in enumerate(attractions,1): add_attraction_card(doc,item,images.get(attraction_id(item)) or images.get(item.get('name')),root,test_mode,index,args.image_root)
    add_section_heading(doc,'餐饮候选与排队备选',1,'section_restaurants',bookmarks['restaurants'])
    if data.get('restaurant_candidates'):
        doc.add_paragraph('以下为带来源的当前参考；个别来源仅取得搜索摘要，详见来源说明。菜单起价和平台人均不是全餐报价；预约、坐席、排队和目标日期营业仍按标注复核。')
        for candidate in data['restaurant_candidates']:
            if not candidate.get('name') or not candidate.get('address'): continue
            doc.add_heading(str(candidate['name']),2)
            doc.add_paragraph('地址：'+str(candidate['address'])+'；餐饮类型：'+public_fact(candidate.get('cuisine','待核验')))
            if candidate.get('cuisine')=='たこ焼': doc.add_paragraph('这是章鱼烧小吃；若想吃完整一餐，优先使用下列拉面备选。坐席与排队情况仍待核实。')
            doc.add_paragraph('营业：'+str(candidate.get('opening_hours','待核验'))+'；价格参考：'+str(candidate.get('price_reference','尚未获取')))
            doc.add_paragraph('预约：'+str(candidate.get('reservation','未取得明确预约条件')))
            if candidate.get('backup'):
                doc.add_paragraph('备选：'+str(candidate['backup'])+'；地址：'+str(candidate.get('backup_address','待核验'))+'；营业：'+str(candidate.get('backup_opening','待核验'))+'；餐饮类型：'+public_fact(candidate.get('backup_cuisine','待核验')))
                if candidate.get('backup_price_reference'): doc.add_paragraph('备选菜单参考：'+candidate['backup_price_reference']+'；连锁起价，具体门店可能不同。')
            sources={e['source_url'] for e in candidate.get('field_evidence',{}).values()}
            for url in sorted(sources): add_hyperlink(doc.add_paragraph(),'打开餐厅来源',url)
    add_section_heading(doc,'每日路线安排',1,'section_daily',bookmarks['daily'])
    for index,day in enumerate(data.get('days',[]),1):
        doc.add_heading(day_heading(index,day.get('date')),2)
        if data.get('reference_budget'):
            times=[str(n.get('time')) for n in day.get('execution_graph',{}).get('nodes',[]) if n.get('time') and n.get('role')=='ACTIVITY']
            departure=times[0] if times else '按当天进度，不早于旅行画像中的作息限制'
            doc.add_paragraph('建议出发：'+str(day.get('departure_suggestion') or (departure+'（规划参考，不为黄金时段自行提前）' if day.get('day_type') not in {'ARRIVAL','DEPARTURE'} else '随核验后的目标日期往返交通安排')))
            meal_nodes=[n for n in day.get('execution_graph',{}).get('nodes',[]) if n.get('role')=='MEAL']
            for n in meal_nodes:
                c=next((c for c in data.get('restaurant_candidates',[]) if c.get('place_id')==n['place_id']),{})
                doc.add_paragraph('餐饮主选：'+str(n['place'])+'；用餐时间：'+str(n.get('time') or '随当天进度调整')+'；营业参考：'+str(c.get('opening_hours') or '待核验')+'；单项菜单参考：'+str(c.get('price_reference') or '未取得'))
                if c.get('backup'):doc.add_paragraph('排队备选：'+str(c['backup'])+'。备选不作为必须额外到访的站点。')
            walk_limit=data.get('trip_profile',{}).get('walking_tolerance_steps')
            if walk_limit:doc.add_paragraph(f'步行上限：每天最多{walk_limit:,}步。全天距离资料仍不完整，当前不能保证合规；保留休息，现场以步数决定缩短户外游览。')
            for c in data.get('research_decision_ledger',[]):
                if c.get('date')==day['date'] and c['decision_type']=='RESTAURANT':
                    doc.add_paragraph('为什么选这家：公开地址与住宿地址一致，午间营业；晚间才开的候选已排除。具体楼层通道、排队和完整餐费仍待核验。')
            for rev in data.get('decision_revision_ledger',[]):
                if rev.get('date')==day['date'] and rev.get('trigger')=='RESEARCH_DRIVEN_DECISION_CHANGE':
                    route=rev.get('new_decision',{}).get('route',{});known=route.get('known_component_minutes')
                    labels={'WALK':'步行','TRAIN':'铁路组合','METRO':'地铁'}
                    old_mode=labels.get(rev['old_decision'].get('mode'),'原方式');new_mode=labels.get(route.get('mode'),'新方式')
                    doc.add_paragraph(f'路线调整：{old_mode}暂改为{new_mode}；已知组件合计约{known}分钟。综合比较步行负荷、换乘和不确定性后选择；完整连接时间与全天步数仍待核，疲劳或下雨时使用交通替代。')
        closure=day.get('route_closure',{})
        if args.technical_debug_report:
            doc.add_paragraph(f"当日路线：{closure.get('resolved_edges',0)}/{closure.get('required_edges',0)} 段有参考或核验数据；覆盖率 {closure.get('route_coverage',0):.0%}；核验状态 {public_status(closure.get('route_verification_status','UNAVAILABLE'))}。")
        p=doc.add_paragraph(); p.add_run('主题：').bold=True; p.add_run(str(day.get('theme',''))+'　')
        p.add_run('区域：').bold=True; p.add_run(str(day.get('area','待确认'))+'　')
        p.add_run('体力：').bold=True; p.add_run(str(day.get('effort','中'))+'　')
        p.add_run('预计步行：').bold=True; p.add_run(str(day.get('walking_estimate','估算'))+'　')
        p.add_run('预计消费：').bold=True; p.add_run(str(day.get('estimated_cost','估算')))
        if day.get('route_chain') or day.get('route_summary'):
            p=doc.add_paragraph(); p.add_run('路线建议：').bold=True; p.add_run(route_summary(day,technical=args.technical_debug_report))
        add_route_table(doc,day,technical=args.technical_debug_report)
        table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
        for cell,label in zip(table.rows[0].cells,('时间','活动','安排提示/费用')): cell.text=label
        activities=day.get('activities') or [{'time':n.get('time','随当天进度'),'activity':n.get('place',''),'note':n.get('note') or n.get('notes','规划建议；实际时刻和接驳待核')} for n in day.get('execution_graph',{}).get('nodes',[])]
        for activity in activities:
            row=table.add_row().cells; row[0].text=str(activity.get('time') or '随当天进度'); row[1].text=public_place(activity.get('activity') or activity.get('place') or ''); row[2].text=str(activity.get('note') or activity.get('notes') or '规划建议；具体条件见景点和交通说明')
        if not activities and day.get('day_type') in {'ARRIVAL','DEPARTURE'}:
            row=table.add_row().cells; row[0].text='随目标日期往返交通核验'
            row[1].text='往返交通与接驳，见上述移动表'
            row[2].text='具体时刻、办理手续和缓冲时间尚未取得；不安排额外景点。'
        format_table(table,header_rows=1)
        if day.get('plan_b_route'):
            mapped='；'.join(str(stop.get('alternative') or stop.get('name') or '') for stop in day['plan_b_route'] if isinstance(stop,dict))
            doc.add_paragraph('天气/变更备选：'+(mapped or '优先替换同一区域的户外活动，保留当天路线结构。'))
        elif day.get('rain_plan'): doc.add_paragraph('天气/变更备选：'+str(day['rain_plan']))
        if day.get('risk'): doc.add_paragraph('注意：'+str(day['risk']))
    omitted=data.get('omitted_attractions',[])
    if omitted:
        add_section_heading(doc,'本次主动舍弃',1,'section_omitted',bookmarks['omitted'])
        for item in omitted:
            if isinstance(item,dict): doc.add_paragraph(f"{item.get('name','候选景点')}：{item.get('why_not','不符合本次旅行重点')} 适合：{item.get('best_for','')}",style='List Bullet')
            else: doc.add_paragraph(str(item),style='List Bullet')
    add_section_heading(doc,'省钱策略与 Plan B',1,'section_alternatives',bookmarks['alternatives'])
    for field in ('savings','food_recommendations','risks'):
        values=data.get(field,[])
        if isinstance(values,str): values=[values]
        for value in values: doc.add_paragraph(str(value),style='List Bullet')
    add_section_heading(doc,'预约与行前清单',1,'section_checklist',bookmarks['checklist'])
    for field in ('bookings','checklist'):
        values=data.get(field,[])
        if isinstance(values,str): values=[values]
        for value in values: doc.add_paragraph(str(value),style='List Bullet')
    add_section_heading(doc,'来源索引',1,'section_sources',bookmarks['sources'])
    table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
    for cell,label in zip(table.rows[0].cells,('来源','链接','核验时间')): cell.text=label
    for source in public_sources(data):
        row=table.add_row().cells; row[0].text=source['title']; row[1].text=''; add_hyperlink(row[1].paragraphs[0],'打开来源',source['url']); row[2].text=source['checked_at'] or '未记录'
    format_table(table,header_rows=1)
    footer=section.footer.paragraphs[0]; footer.alignment=2; footer.add_run('第 ')
    from docx.oxml import OxmlElement
    run=OxmlElement('w:r'); begin=OxmlElement('w:fldChar'); begin.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType','begin'); run.append(begin); footer._p.append(run)
    run=OxmlElement('w:r'); instr=OxmlElement('w:instrText'); instr.text=' PAGE '; run.append(instr); footer._p.append(run)
    run=OxmlElement('w:r'); end=OxmlElement('w:fldChar'); end.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType','end'); run.append(end); footer._p.append(run); footer.add_run(' 页')
    order_handbook_sections(doc,sections)
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output); print(output)


if __name__=='__main__': main()
