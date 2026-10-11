"""Release/presentation contracts; not real client installation or live research."""
import copy,hashlib,importlib.util,json,subprocess,sys,tempfile,unittest,zipfile
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from build_docx import day_heading,route_summary,add_route_table,add_budget_cards
from package_plugin import build,distribution_files

def module(name,path):
    spec=importlib.util.spec_from_file_location(name,ROOT/path);result=importlib.util.module_from_spec(spec);spec.loader.exec_module(result);return result

release=module('verify_release_contract','installer/verify_release.py')
structure=module('distribution_contract','installer/validate_distribution.py')

class Presentation(unittest.TestCase):
    def test_day_relative_title_once(self):
        self.assertEqual(day_heading(1,'Day 1（日期未定）'),'Day 1（日期未定）')
        self.assertEqual(day_heading(2,'Day 2'),'Day 2（日期未定）')
    def test_real_date_preserved(self):self.assertEqual(day_heading(2,'2027-04-21'),'Day 2 · 2027-04-21')
    def test_unresolved_edges_have_natural_explanation(self):
        day={'route_chain':[{'from':'AIRPORT_UNRESOLVED','to':'AIRPORT_UNRESOLVED'}]}
        text=route_summary(day);self.assertNotIn('→',text);self.assertIn('机场',text)
    def test_real_places_without_fake_time(self):
        text=route_summary({'route_chain':[{'from':'Museum','to':'Park','duration_minutes':None}]})
        self.assertEqual(text,'Museum → Park');self.assertNotIn('分钟',text)
    def test_source_graph_not_mutated_by_presentation(self):
        day={'route_chain':[{'from':'HOTEL_UNRESOLVED','to':'Park'}]};before=copy.deepcopy(day)
        route_summary(day);self.assertEqual(day,before)
    def test_unknown_route_table_omitted(self):
        from docx import Document
        doc=Document();add_route_table(doc,{'route_chain':[{'from':'Museum','to':'Park'}]})
        self.assertEqual(len(doc.tables),0)
    def test_known_route_still_present_and_qualified(self):
        from docx import Document
        doc=Document();add_route_table(doc,{'route_chain':[{'from':'Museum','to':'Park','duration_minutes':20,'status':'ESTIMATED'}]})
        text=' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('20',text);self.assertIn('估算',text)
    def test_technical_mode_retains_route_diagnostic(self):
        self.assertIn('待核验',route_summary({'route_chain':[{'from':'Museum','to':'Park'}]},technical=True))
    def test_unknown_costs_not_spread_across_duplicate_rows(self):
        from docx import Document
        doc=Document();data={'budget_reconciliation':{'currency':'CNY','user_budget_limit_party':6000,
             'unknown_cost_categories':['flight','hotel'],'rows':[{'category':'flight','amount':None}]}}
        add_budget_cards(doc,data);self.assertEqual(len(doc.tables),1)
        text=' '.join(p.text for p in doc.paragraphs);self.assertIn('机票',text);self.assertIn('判断',text)

class Distribution(unittest.TestCase):
    def test_actual_marketplace_structure(self):self.assertEqual(structure.validate(ROOT)['structural_status'],'PASS')
    def test_manifest_is_not_host_pass(self):
        result=structure.validate(ROOT)
        for key in ('claude_persistent_installation','codex_gui_installation','doubao_host_installation'):self.assertEqual(result[key],'NOT_TESTED')
    def test_source_path_from_root(self):
        catalog=json.loads((ROOT/'.claude-plugin/marketplace.json').read_text())
        self.assertEqual((ROOT/catalog['plugins'][0]['source']).resolve(),ROOT)
    def test_catalog_identity_in_guides(self):
        for name in ('README.md','README-BEGINNER.md','README-ADVANCED.md'):
            self.assertIn('smart-travel-planner-marketplace',(ROOT/name).read_text(encoding='utf-8'))
    def test_declared_repository_matches_guide(self):
        repository=json.loads((ROOT/'config/release.json').read_text())['repository']
        guide=(ROOT/'README-BEGINNER.md').read_text(encoding='utf-8')
        if repository is None:
            self.assertIn('占位',guide)
        else:
            self.assertRegex(repository,r'^[A-Za-z0-9][A-Za-z0-9_.-]*/[A-Za-z0-9][A-Za-z0-9_.-]*$')
            self.assertIn(repository+'#v1.0.2',guide)
            self.assertNotIn('YOUR_GITHUB_OWNER/YOUR_REPOSITORY',guide)
    def test_no_user_work_or_unapproved_image(self):
        names=[p.relative_to(ROOT).as_posix() for p in distribution_files()]
        self.assertFalse(any(n.startswith(('work/','outputs/','acceptance/')) or '__pycache__' in n or n=='.env' for n in names))
        test_images=[n for n in names if n.startswith('tests/') and n.endswith(('.jpg','.png'))]
        self.assertEqual(test_images,['tests/fixtures/images/sensoji-cc0.jpg'])
    def test_zip_checksum_matches(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=build(Path(tmp)/'release.zip');result=release.verify(path,str(path)+'.sha256')
            self.assertEqual(result['version'],json.loads((ROOT/'plugin.json').read_text())['version'])
            self.assertFalse(result['execution_performed'])
    def test_tampered_zip_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=build(Path(tmp)/'release.zip');path.write_bytes(path.read_bytes()+b'changed')
            with self.assertRaisesRegex(ValueError,'SHA256 mismatch'):release.verify(path,str(path)+'.sha256')
    def test_archive_traversal_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'release.zip'
            with zipfile.ZipFile(path,'w') as z:z.writestr('../escape.py','no execution')
            side=Path(str(path)+'.sha256');side.write_text(hashlib.sha256(path.read_bytes()).hexdigest()+'  '+path.name)
            with self.assertRaisesRegex(ValueError,'Unsafe archive path'):release.verify(path,side)
    def test_ai_guide_permission_and_degradation(self):
        text=(ROOT/'README-AI-INSTALL.md').read_text(encoding='utf-8')
        for required in ('批准','SHA256','NOT_TESTED','不保证','完整插件','不会安装'):self.assertIn(required,text)
    def test_github_template_no_network_mutation(self):
        from prepare_github_release import prepare
        with tempfile.TemporaryDirectory() as tmp:
            result=prepare(Path(tmp)/'export')
            self.assertIsNone(result['repository']);self.assertFalse(result['upload_performed'])
            self.assertTrue((Path(result['github_ready_directory'])/'.claude-plugin/marketplace.json').is_file())
    def test_bad_repository_target_rejected(self):
        from prepare_github_release import prepare
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):prepare(Path(tmp)/'export','https://github.com/untrusted/repo;command')

if __name__=='__main__':unittest.main()
