"""Synthetic regression contracts only; not the three real release cases."""
import sys,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from test_autonomous_research import minimal,gap,NOW
from semantic_query_planner import semantic_query
from execution_readiness import evaluate_plan
from research_gap_detector import detect_gaps
from cost_ledger import reconcile_budget
from reference_budget_estimator import estimate_budget
from build_docx import add_section_heading,order_handbook_sections,public_cost

def rail_plan():
    p=minimal();p['trip_profile'].update(country='China',destination='杭州',destination_currency='CNY',research_language='zh',
        domestic_trip=True,intercity_mode='TRAIN',airfare_included=False,hotel='经济酒店',hotel_address='中河中路175号',
        origin_station='上海虹桥站',arrival_station='杭州东站',departure_station='杭州东站',return_station='上海虹桥站',
        lodging_area='湖滨—解放路')
    p['days']=[{'date':'2027-04-20','day_type':'ARRIVAL'}, {'date':'2027-04-21','day_type':'DEPARTURE'}]
    return p

class ReleaseRegressions(unittest.TestCase):
    def test_rail_anchors_survive(self):
        p=evaluate_plan(rail_plan(),now=NOW)
        self.assertEqual(p['days'][0]['execution_graph']['nodes'][0]['role'],'STATION')
        self.assertEqual(p['days'][-1]['execution_graph']['nodes'][-1]['place'],'上海虹桥站')
    def test_rail_not_bookable_without_tickets(self):
        p=evaluate_plan(rail_plan(),now=NOW)
        self.assertEqual(p['readiness']['booking_readiness'],'NOT_READY')
        self.assertIn('INTERCITY_TICKET_UNVERIFIED',p['readiness']['blocking_reasons'])
    def test_rail_no_false_airport_requirement(self):
        p=evaluate_plan(rail_plan(),now=NOW)
        self.assertNotIn('AIRPORT_UNRESOLVED',p['readiness']['blocking_reasons'])
        self.assertEqual(p['state_consistency']['status'],'PASS')
    def test_domestic_no_japan_research(self):
        p=evaluate_plan(rail_plan(),now=NOW);g=detect_gaps(p,now=NOW)
        self.assertFalse(any(x['target']['kind'] in {'entry','exchange','flight','flight_static'} for x in g))
    def test_domestic_missing_costs_exclude_flight_visa(self):
        p=evaluate_plan(rail_plan(),now=NOW)
        self.assertNotIn('flight',p['budget_reconciliation']['unknown_cost_categories'])
        self.assertNotIn('visa_related',p['budget_reconciliation']['unknown_cost_categories'])
    def test_domestic_budget_stays_unknown(self):
        p=evaluate_plan(rail_plan(),now=NOW)
        self.assertIsNone(p['budget_reconciliation']['per_person_total'])
        self.assertEqual(p['budget_reconciliation']['feasibility'],'UNVERIFIED')
    def test_chinese_restaurant_discovery(self):
        p=rail_plan();g=gap('RESTAURANT_STATIC','restaurant','name','酒店附近晚餐');g['target'].update(area='湖滨—解放路')
        q=semantic_query([g],p['trip_profile'],plan=p)
        self.assertEqual(q['quality_gate']['status'],'PASS');self.assertEqual(q['language'],'zh')
        self.assertNotIn('Japan',q['query']);self.assertIn('餐厅',q['query'])
    def test_station_pair_not_walk_query(self):
        p=evaluate_plan(rail_plan(),now=NOW);g=next(g for g in detect_gaps(p,now=NOW) if g['target'].get('from_place_id')=='origin-station')
        q=semantic_query([g],p['trip_profile'],plan=p)
        self.assertEqual(q['intent']['information_need'],'INTERCITY_RAIL');self.assertNotIn('步行',q['query'])
    def test_rail_price_coverage_not_japan_categories(self):
        p=evaluate_plan(rail_plan(),now=NOW);r=estimate_budget(p,now=NOW)
        self.assertNotIn('flight',r['cost_reference_coverage']['required_categories'])
    def test_handbook_section_move_keeps_content(self):
        from docx import Document
        d=Document();d.add_paragraph('封面');add_section_heading(d,'资料',1,'section_readiness',1);d.add_paragraph('资料正文')
        add_section_heading(d,'行程',1,'section_daily',2);d.add_paragraph('行程正文')
        order_handbook_sections(d,[('行程','daily'),('资料','readiness')])
        self.assertEqual([p.text for p in d.paragraphs],['封面','行程','行程正文','资料','资料正文'])
    def test_unknown_cost_not_displayed_as_free(self):
        self.assertNotIn('免费',public_cost(0,'CNY',unknown=True))
        self.assertNotIn('0.00',public_cost(0,'CNY',unknown=True))

if __name__=='__main__':unittest.main()
